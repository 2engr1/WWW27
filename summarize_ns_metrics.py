#!/usr/bin/env python3
"""Only collect precomputed averages from five NS modes across six scales."""
import argparse
import csv
import json
from pathlib import Path

SCALES = [('20K','20000'), ('50K','50000'), ('100K','100000'),
          ('200K','200000'), ('400K','400000'), ('Full','full')]
MODES = [('ns-retrieval','NS-Retrieval'), ('ns-ce','NS+CE'),
         ('ns','NS'), ('ns-nec','NS-Nec'), ('ns-suf','NS-Suf')]
METRICS = [('MRR@5','mrr@5'), ('Recall@5','recall@5'),
           ('MRR@10','mrr@10'), ('Recall@10','recall@10'),
           ('MRR@20','mrr@20'), ('Recall@20','recall@20'),
           ('MRR@50','mrr@50'), ('Recall@50','recall@50')]
HEADER = ['Scale'] + [name for name, _ in METRICS]

def md_table(rows):
    return '\n'.join(['| ' + ' | '.join(HEADER) + ' |',
                      '| ' + ' | '.join(['---'] * len(HEADER)) + ' |'] +
                     ['| ' + ' | '.join(row) + ' |' for row in rows]) + '\n'

def main():
    parser = argparse.ArgumentParser(description='Summarize existing JSON metrics. No metric recomputation.')
    parser.add_argument('--root', default='outputs')
    parser.add_argument('--out', default=None)
    parser.add_argument('--digits', type=int, default=4)
    parser.add_argument('--strict', action='store_true', help='Return nonzero if a result is missing')
    args = parser.parse_args()
    root = Path(args.root).expanduser()
    out = Path(args.out).expanduser() if args.out else root / 'summary_tables'
    out.mkdir(parents=True, exist_ok=True)
    warnings, combined, md_sections = [], [], []
    for mode, display in MODES:
        rows = []
        for scale, dirname in SCALES:
            path = root / dirname / f'{mode}.jsonl.summary.json'
            avg = {}
            if not path.is_file():
                warnings.append(f'MISSING FILE: {path}')
            else:
                try:
                    data = json.loads(path.read_text(encoding='utf-8'))
                    avg = data['averages']
                    if not isinstance(avg, dict):
                        raise ValueError('averages is not an object')
                except (OSError, json.JSONDecodeError, KeyError, ValueError) as exc:
                    warnings.append(f'INVALID FILE: {path}: {exc}')
                    avg = {}
            vals = []
            for _, metric in METRICS:
                value = avg.get(metric)
                try:
                    if value is None or isinstance(value, bool):
                        raise ValueError('missing/non-numeric')
                    vals.append(f'{float(value):.{args.digits}f}')
                except (ValueError, TypeError, OverflowError):
                    vals.append('')
                    warnings.append(f'MISSING OR INVALID: {path} -> averages.{metric}')
            row = [scale, *vals]
            rows.append(row)
            combined.append([display, *row])
        with (out / f'{mode}_table.csv').open('w', encoding='utf-8-sig', newline='') as f:
            writer = csv.writer(f); writer.writerow(HEADER); writer.writerows(rows)
        md_sections.append(f'## {display}\n\n{md_table(rows)}')
    with (out / 'all_modes.csv').open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.writer(f); writer.writerow(['Mode', *HEADER]); writer.writerows(combined)
    (out / 'all_modes.md').write_text('\n'.join(md_sections), encoding='utf-8')
    (out / 'warnings.txt').write_text('\n'.join(warnings) + ('\n' if warnings else 'No missing files or metrics.\n'), encoding='utf-8')
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        all_data = [('All Modes', ['Mode', *HEADER], combined)]
        for mode, display in MODES:
            with (out / f'{mode}_table.csv').open('r', encoding='utf-8-sig', newline='') as f:
                matrix = list(csv.reader(f))
            all_data.append((display, matrix[0], matrix[1:]))
        for sheet_name, head, rows in all_data:
            ws = wb.create_sheet(sheet_name)
            ws.append(head)
            for row in rows:
                converted = [float(v) if i >= (2 if sheet_name == 'All Modes' else 1) and v else v for i, v in enumerate(row)]
                ws.append(converted)
            ws.freeze_panes = 'C2' if sheet_name == 'All Modes' else 'B2'
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='254B70')
                cell.alignment = Alignment(horizontal='center')
            for col in range(1, len(head) + 1):
                ws.column_dimensions[get_column_letter(col)].width = 17 if col <= 2 else 15
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, float): cell.number_format = '0.' + '0'*args.digits
        wb.save(out / 'ns_metrics_tables.xlsx')
    except ImportError:
        warnings.append('openpyxl not installed: skipped Excel; CSV and Markdown are available')
        (out / 'warnings.txt').write_text('\n'.join(warnings) + '\n', encoding='utf-8')
    print(f'Collected {len(combined)} mode-scale rows into {out}')
    print('\n'.join(md_sections))
    if warnings:
        print(f'{len(warnings)} warning(s), see {out / "warnings.txt"}')
        if args.strict: raise SystemExit(1)

if __name__ == '__main__':
    main()
